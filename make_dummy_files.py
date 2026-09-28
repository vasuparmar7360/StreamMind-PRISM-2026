from fpdf import FPDF
import docx

# Make PDF
pdf = FPDF()
pdf.add_page()
pdf.set_font("Arial", size=12)
pdf.cell(200, 10, txt="This is a test PDF document for OwnMind AI.", ln=1, align='L')
pdf.cell(200, 10, txt="It has multiple pages and text.", ln=1, align='L')
pdf.add_page()
pdf.cell(200, 10, txt="This is page 2.", ln=1, align='L')
pdf.output("test_doc.pdf")

# Make DOCX
doc = docx.Document()
doc.add_paragraph("This is a test DOCX document for OwnMind AI.")
doc.add_paragraph("It contains multiple paragraphs.")
doc.save("test_doc.docx")
